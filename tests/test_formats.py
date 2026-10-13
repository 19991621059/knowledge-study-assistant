import io

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import ArrayObject, DecodedStreamObject, DictionaryObject, NameObject


def test_markdown_and_docx_preserve_text(client):
    text = "# 私人测试资料\n\n这是用户自己的长段文字，导入后需要保留连续内容和定位信息。这个文件不发布为教材。"
    first = client.post("/api/documents", files={"file": ("my.md", text.encode())})
    assert first.status_code == 200
    document = Document()
    document.add_paragraph(text)
    document.add_table(rows=1, cols=1).cell(0, 0).text = "表格里的内容也需要保留"
    document.add_paragraph("表格之后的结尾应当保持在表格之后")
    buffer = io.BytesIO()
    document.save(buffer)
    second = client.post("/api/documents", files={"file": ("my.docx", buffer.getvalue())})
    assert second.status_code == 200
    chunks = client.app.state.store.query(
        "SELECT text FROM chunks WHERE document_id=?", (second.json()["id"],)
    )
    original = "".join(c["text"] for c in chunks)
    assert "表格里的内容也需要保留" in original
    assert original.index("表格里的内容也需要保留") < original.index("表格之后的结尾")


def make_pdf(blank=False):
    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    if not blank:
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        stream = DecodedStreamObject()
        stream.set_data(
            b"BT /F1 12 Tf 40 790 Td (This original document contains searchable text and keeps its physical page number.) Tj ET"
        )
        page[NameObject("/Contents")] = ArrayObject([writer._add_object(stream)])
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def test_pdf_text_page_and_scanned_rejection(client):
    response = client.post("/api/documents", files={"file": ("text.pdf", make_pdf())})
    assert response.status_code == 200, response.text
    chunk = client.app.state.store.query("SELECT * FROM chunks")[0]
    assert chunk["page"] == 1 and "searchable text" in chunk["text"]
    blank = client.post("/api/documents", files={"file": ("scanned.pdf", make_pdf(True))})
    assert blank.status_code == 422 and "OCR" in blank.text


def test_unknown_format_and_invalid_encoding(client):
    assert (
        client.post(
            "/api/documents", files={"file": ("malicious.html", b"<script>alert(1)</script>" * 5)}
        ).status_code
        == 422
    )
    assert (
        client.post("/api/documents", files={"file": ("binary.txt", b"\x00\xff" * 100)}).status_code
        == 422
    )
